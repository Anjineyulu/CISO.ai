"""Generate a Word action plan in memory; no additional model call or storage."""
from io import BytesIO
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core import validate_plan
from session_io import validate_tracking, default_tracking

from governance import current_signoff, default_alignment, CATALOG, NOTICE, ALIGNMENT_NOTICE, NIST_SOURCE, CIS_SOURCE

LOGO = Path(__file__).resolve().parent / 'assets' / 'logo.png'


def text(value):
    # XML 1.0 disallows control characters even if JSON input accepted them.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', str(value))


def shade(element, color):
    fill = OxmlElement('w:shd')
    fill.set(qn('w:fill'), color)
    element.append(fill)


def label_paragraph(doc, label, value):
    p = doc.add_paragraph()
    p.add_run(label + ': ').bold = True
    p.add_run(text(value) or 'Not specified')
    return p


def export_docx(record):
    plan = validate_plan(record['reviewed_plan'], record['input'])
    tracking = validate_tracking(record.get('tracking', default_tracking(plan)), plan)
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].font.name = 'Calibri'
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Heading 1'].font.size = Pt(18)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
    doc.core_properties.title = 'Security remediation action plan'
    doc.core_properties.author = 'CISO.ai'
    doc.core_properties.subject = 'Human reviewed remediation actions'
    if LOGO.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        shade(p._p.get_or_add_pPr(), '13263C')
        run = p.add_run()
        shape = run.add_picture(str(LOGO), width=Inches(3.25))
        shape._inline.docPr.set('descr', 'CISO.ai Security Operating System logo')
        p.paragraph_format.space_after = Pt(14)
    doc.add_paragraph('Security remediation action plan', 'Title')
    doc.add_paragraph('This plan records the actions reviewed by the user for the findings below. Confirm change approval and operational impact before implementation. Review and a reported completion status do not certify security or verify that remediation has occurred.')
    label_paragraph(doc, 'Model', record['model'])
    label_paragraph(doc, 'Draft generated', record['generated_at'])
    label_paragraph(doc, 'Review recorded', record['reviewed_at'])
    doc.add_heading('Business context', level=1)
    doc.add_paragraph(text(record['input']['business_context']))
    label_paragraph(doc, 'Operational constraints', record['input'].get('constraints', ''))
    doc.add_heading('CISO decision', level=1)
    doc.add_paragraph(NOTICE)
    signoff = current_signoff(record)
    if signoff:
        for key, label in [('decision', 'Decision'), ('name', 'Reviewer'), ('role', 'Role'), ('rationale', 'Rationale and conditions'), ('signed_at', 'Recorded at'), ('record_digest', 'Reviewed record SHA-256')]:
            label_paragraph(doc, label, signoff[key])
    else:
        doc.add_paragraph('Not signed off')
    doc.add_heading('Standards alignment scope', level=1)
    doc.add_paragraph(ALIGNMENT_NOTICE)
    doc.add_paragraph('Sources: ' + NIST_SOURCE + ' ; ' + CIS_SOURCE)
    doc.add_heading('Action overview', level=1)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [.6, 1.12, 1.9, 1.3, 1.9]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    headers = ['Finding', 'Priority', 'Owner', 'Target date', 'Status']
    for cell, heading in zip(table.rows[0].cells, headers):
        cell.text = heading
    trpr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    trpr.append(repeat)
    for item in plan['items']:
        fid = item['finding_id']
        track = tracking[fid]
        vals = [fid, item['priority'], track['owner'] or 'Unassigned', track['target_date'] or 'Not set', track['status']]
        cells = table.add_row().cells
        for cell, value in zip(cells, vals):
            cell.text = text(value)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade(props, '173459' if i == 0 else ('F1F5F9' if i % 2 else 'FFFFFF'))
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                border = OxmlElement('w:' + edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    border.set(qn('w:' + key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), '90')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 2 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '000000')
    sources = {row['finding_id']: row for row in record['input']['findings']}
    for item in plan['items']:
        fid = item['finding_id']
        source, track = sources[fid], tracking[fid]
        doc.add_page_break()
        doc.add_heading('Finding ' + fid, level=1)
        label_paragraph(doc, 'Asset alias', source['asset'])
        label_paragraph(doc, 'Priority', item['priority'])
        label_paragraph(doc, 'Owner role or alias', track['owner'] or 'Unassigned')
        label_paragraph(doc, 'Target date', track['target_date'] or 'Not set')
        label_paragraph(doc, 'Reported action status', track['status'])
        doc.add_heading('Observed issue', level=2)
        doc.add_paragraph(text(source['observation']))
        label_paragraph(doc, 'Reported severity', source['severity'])
        label_paragraph(doc, 'Exposure', source['exposure'])
        label_paragraph(doc, 'Business criticality', source['criticality'])
        for field, heading in [('rationale', 'Priority rationale'), ('action', 'Reviewed action'),
                               ('verification', 'Verification steps'), ('uncertainty', 'Assumptions and unknowns')]:
            doc.add_heading(heading, level=2)
            doc.add_paragraph(text(item[field]))
        mapping = record.get('alignment', default_alignment(plan))[fid]
        doc.add_heading('Standards alignment · advisory', level=2)
        label_paragraph(doc, 'References', '; '.join(f'{ref} — {CATALOG[ref]}' for ref in mapping['references']) or 'None selected')
        label_paragraph(doc, 'Assessment', mapping['assessment'])
        label_paragraph(doc, 'Mapping rationale', mapping['rationale'])
        label_paragraph(doc, 'Evidence reference', mapping['evidence'])
    # A short page number supports longer reports without adding identifying data.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('CISO.ai  |  Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
