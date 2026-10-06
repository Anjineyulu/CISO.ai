import io
import unittest
import zipfile
from docx import Document
from word_export import export_docx
from test_session_io import record_fixture
from test_core import plan_fixture


class WordExportTests(unittest.TestCase):
    def test_export_contains_reviewed_action_tracking_and_logo(self):
        record = record_fixture()
        plan = plan_fixture()
        plan['items'][0]['action'] = 'Human edited action for this test.'
        record.update(reviewed_plan=plan, reviewed_at='2026-10-04T12:00:00Z',
                      tracking={'F1': dict(owner='Test owner', target_date='2026-11-01', status='In progress')})
        data = export_docx(record)
        doc = Document(io.BytesIO(data))
        contents = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Human edited action for this test.', contents)
        self.assertIn('Test owner', contents)
        self.assertIn('2026-11-01', contents)
        self.assertIn('In progress', contents)
        self.assertEqual(len(doc.tables[0].rows), 2)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.assertTrue(any(n.startswith('word/media/') for n in z.namelist()))
            self.assertNotIn(b'GEMINI_API_KEY', z.read('word/document.xml'))

    def test_control_characters_are_removed(self):
        record = record_fixture()
        plan = plan_fixture()
        plan['items'][0]['action'] = 'Test action\x00 with XML-invalid character.'
        record.update(reviewed_plan=plan, reviewed_at='2026-10-04T12:00:00Z')
        data = export_docx(record)
        doc = Document(io.BytesIO(data))
        self.assertIn('Test action with XML-invalid character.', '\n'.join(p.text for p in doc.paragraphs))
